"""Optional document text extraction (PDF/DOCX/OCR) and model inference."""

import os
import json
import re
from functools import lru_cache
from pathlib import Path

MAX_TEXT_CHARS = 100_000
MAX_PDF_PAGES = 50
MAX_OCR_PDF_PAGES = 10
OCR_TIMEOUT_SECONDS = 20


def extract_document_entities(text):
    """Extract conservative, explicitly labelled fields as unverified drafts."""
    if not text:
        return {}

    normalized = re.sub(r"[ \t]+", " ", text)
    patterns = {
        "osca_id": r"\bOSCA\s*(?:I\s*\.?\s*D\s*\.?|ID)\s*(?:NO\.?|NUMBER|#)?\s*[:#\-]?\s*([A-Z0-9][A-Z0-9\-/]{2,})",
        "registration_cert_no": r"\b(?:REGISTRATION\s+CERTIFICATE|REG\.?\s*CERT\.?|CERTIFICATE)\s*(?:NO\.?|NUMBER|#)\s*[:#\-]?\s*([A-Z0-9][A-Z0-9\-/]{2,})",
        "date_of_birth": r"\b(?:DATE\s+OF\s+BIRTH|BIRTH\s+DATE)\s*[:#\-]?\s*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}-\d{2}-\d{2})",
        "last_name": r"\bLAST\s+NAME\s*[:\-]\s*([A-Z][A-Z' -]{1,60})",
        "first_name": r"\b(?:FIRST|GIVEN)\s+NAME\s*[:\-]\s*([A-Z][A-Z' -]{1,60})",
        "middle_name": r"\bMIDDLE\s+NAME\s*[:\-]\s*([A-Z][A-Z' -]{1,60})",
    }
    entities = {}
    for field, pattern in patterns.items():
        match = re.search(pattern, normalized, re.IGNORECASE)
        if match:
            value = re.split(r"\b(?:AGE|SEX|ADDRESS|BARANGAY|CITY|PROVINCE|CIVIL STATUS)\b", match.group(1), maxsplit=1, flags=re.IGNORECASE)[0]
            value = value.strip(" .,:;-/")
            if value:
                entities[field] = value[:100]
    return entities


def classifier_artifact_path():
    use_demo = os.environ.get("ARKHEION_USE_DEMO_CLASSIFIER", "").strip().lower() in {"1", "true", "yes"}
    configured = os.environ.get("ARKHEION_DOCUMENT_CLASSIFIER_PATH")
    if configured:
        configured_path = Path(configured)
        # Ignore a stale demo-model path unless demo mode is explicitly enabled.
        if use_demo or configured_path.name != "simulated-demo.joblib":
            return configured_path
    from django.conf import settings

    model_dir = Path(settings.BASE_DIR) / "var"
    if use_demo:
        # Development only: this artifact was trained on simulated examples.
        return model_dir / "simulated-demo.joblib"
    # Production mode is fail-closed: no model suggestion until an evaluated
    # artifact has been trained from an approved, anonymized labeled dataset.
    return model_dir / "document_classifier.joblib"


def classifier_source(artifact=None):
    """Describe model provenance without presenting a custom model as evaluated."""
    artifact = artifact or classifier_artifact_path()
    if not artifact.is_file():
        return ""
    use_demo = os.environ.get("ARKHEION_USE_DEMO_CLASSIFIER", "").strip().lower() in {"1", "true", "yes"}
    if use_demo or "simulated" in artifact.name.lower():
        return "demo"
    manifest_path = artifact.with_suffix(artifact.suffix + ".evaluation.json")
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        if manifest.get("evaluation_type") == "independent_test" and manifest.get("passed") is True:
            return "evaluated" if manifest.get("dataset_kind") == "anonymized_osca" else "simulated_test"
    except (OSError, ValueError, TypeError):
        pass
    return "configured"


def _configure_ocr():
    import pytesseract

    executable = os.environ.get("TESSERACT_CMD")
    if executable:
        pytesseract.pytesseract.tesseract_cmd = executable
    return pytesseract


@lru_cache(maxsize=1)
def _rapid_ocr_engine():
    """Load the bundled CPU OCR engine once per server process."""
    from rapidocr import RapidOCR

    return RapidOCR()


def _ocr_with_rapid(image):
    """OCR using RapidOCR, which does not need a separate Windows install."""
    try:
        result = _rapid_ocr_engine()(image)
        # RapidOCR returns an object with `txts` in current releases.
        words = getattr(result, "txts", None)
        if words is None and isinstance(result, tuple):
            words = result[0]
        text = "\n".join(str(word).strip() for word in (words or []) if str(word).strip())
        return text, "ocr_complete" if text else "no_readable_text"
    except ImportError:
        return "", "ocr_dependencies_missing"
    except Exception:
        return "", "ocr_error"


def _ocr_image(image):
    """Return (text, status) for a PIL image, without making upload fail."""
    try:
        from PIL import ImageOps
        image = ImageOps.exif_transpose(image).convert("RGB")
        image.thumbnail((2400, 2400))
        try:
            pytesseract = _configure_ocr()
            executable = os.environ.get("TESSERACT_CMD")
            if not executable or Path(executable).is_file():
                text = pytesseract.image_to_string(
                    image,
                    lang=os.environ.get("ARKHEION_OCR_LANGUAGE", "eng"),
                    config="--psm 6",
                    timeout=OCR_TIMEOUT_SECONDS,
                ).strip()
                if text:
                    return text, "ocr_complete"
        except Exception:
            # Tesseract is optional; use RapidOCR if it's not installed or fails.
            pass
        try:
            return _ocr_with_rapid(image)
        except Exception:
            return "", "ocr_error"
    except ImportError:
        return "", "ocr_dependencies_missing"
    except Exception:
        return "", "ocr_error"


def _ocr_pdf(source):
    """OCR up to ten pages of a scanned PDF using the configured OCR engine."""
    try:
        import fitz
        from PIL import Image
    except ImportError:
        return "", "ocr_dependencies_missing"

    try:
        source.seek(0)
        pdf = fitz.open(stream=source.read(), filetype="pdf")
        page_text = []
        statuses = []
        for page_index in range(min(pdf.page_count, MAX_OCR_PDF_PAGES)):
            page = pdf.load_page(page_index)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(1.7, 1.7), alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            text, status = _ocr_image(image)
            if text:
                page_text.append(text)
            statuses.append(status)
        text = "\n".join(page_text).strip()[:MAX_TEXT_CHARS]
        if text:
            return text, "ocr_complete"
        if "ocr_dependencies_missing" in statuses:
            return "", "ocr_dependencies_missing"
        if "ocr_timeout" in statuses:
            return "", "ocr_timeout"
        return "", "no_readable_text"
    except Exception:
        return "", "ocr_error"


def extract_document_text(document):
    """Return (extracted text, status); never block saving a document on OCR."""
    extension = Path(document.file.name).suffix.lower()
    try:
        with document.file.open("rb") as source:
            if extension == ".pdf":
                try:
                    from pypdf import PdfReader

                    reader = PdfReader(source, strict=False)
                    native_text = "\n".join(
                        (page.extract_text() or "") for page in reader.pages[:MAX_PDF_PAGES]
                    ).strip()
                except ImportError:
                    native_text = ""
                except Exception:
                    native_text = ""
                if len(native_text) >= 40:
                    return native_text[:MAX_TEXT_CHARS], "text_extracted"
                return _ocr_pdf(source)

            if extension in {".png", ".jpg", ".jpeg"}:
                try:
                    from PIL import Image

                    with Image.open(source) as image:
                        return _ocr_image(image)
                except ImportError:
                    return "", "ocr_dependencies_missing"
                except Exception:
                    return "", "ocr_error"

            if extension == ".docx":
                from docx import Document as WordDocument

                word = WordDocument(source)
                parts = [paragraph.text for paragraph in word.paragraphs]
                parts.extend(
                    cell.text
                    for table in word.tables
                    for row in table.rows
                    for cell in row.cells
                )
                text = "\n".join(parts).strip()[:MAX_TEXT_CHARS]
                return text, "text_extracted" if text else "no_readable_text"
    except ImportError:
        return "", "ocr_dependencies_missing"
    except Exception:
        return "", "extraction_error"
    return "", "unsupported_file_type"


def classify_document(document, text=None):
    artifact = classifier_artifact_path()
    if not artifact.is_file():
        return None
    if text is None:
        text, _ = extract_document_text(document)
    if not text:
        return None
    try:
        from joblib import load

        classifier = load(artifact)
        prediction = str(classifier.predict([text])[0])
        confidence = None
        if hasattr(classifier, "predict_proba"):
            probabilities = classifier.predict_proba([text])[0]
            classes = list(classifier.classes_)
            confidence = float(probabilities[classes.index(prediction)])
        return prediction, confidence
    except Exception:
        # Optional ML/model failures must never prevent document upload.
        return None

